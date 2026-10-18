"""Skill: Textextraktion (PDF, Office, Text, Notebooks). Liest Dateien nur lesend."""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

from .filetype import TEXT_EXTS
from .video import VIDEO_EXTS as VIDEO_EXTS_

IMAGE_LIKE = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".heic", ".heif", ".hif", ".avif", ".cr2", ".cr3", ".crw", ".nef",
              ".nrw", ".arw", ".srf", ".sr2", ".dng", ".raf", ".orf", ".rw2", ".pef", ".srw", ".3fr", ".erf", ".kdc", ".mrw",
              ".x3f", ".raw", ".rwl", ".iiq"}
AUDIO_EXTS_ = {".mp3", ".m4a", ".aac", ".flac", ".wav", ".ogg", ".oga", ".opus", ".wma", ".aiff", ".aif", ".alac", ".amr"}

MAX_CHARS = 20000


def _pdf(path: str) -> tuple[str, int]:
    import pymupdf

    with pymupdf.open(path) as doc:
        parts = []
        for page in doc:
            parts.append(page.get_text())
            if sum(map(len, parts)) > MAX_CHARS:
                break
        return "\n".join(parts), doc.page_count


def _pptx(path: str) -> tuple[str, int]:
    from pptx import Presentation

    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"--- Folie {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
            if notes:
                out.append(f"[Notizen] {notes}")
    return "\n".join(out), len(prs.slides)


def _docx(path: str) -> tuple[str, int]:
    import docx

    d = docx.Document(path)
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables[:10]:
        for row in t.rows:
            text += "\n" + " | ".join(c.text for c in row.cells)
    return text, None


def _xlsx(path: str) -> tuple[str, int]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets[:5]:
        out.append(f"--- Blatt {ws.title} ---")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            out.append(" | ".join("" if v is None else str(v) for v in row))
            if i > 60:
                break
    n = len(wb.worksheets)
    wb.close()
    return "\n".join(out), n


def _ipynb(path: str) -> tuple[str, int]:
    nb = json.loads(Path(path).read_text(encoding="utf-8", errors="ignore"))
    cells = nb.get("cells", [])
    return "\n".join("".join(c.get("source", [])) for c in cells), len(cells)


def _odf(path: str) -> tuple[str, int]:
    import re

    with zipfile.ZipFile(path) as z:
        xml = z.read("content.xml").decode("utf-8", "ignore")
    return re.sub(r"<[^>]+>", " ", xml), None


def _iwork(path: str) -> tuple[str, int | None]:
    """Keynote/Pages/Numbers: Text aus der LibreOffice-PDF (gecacht, wird auch für die Miniaturen genutzt)."""
    from . import formats, thumbnails

    if not formats.is_iwork(path):
        return "", None  # z. B. Schlüsseldatei „server.key“ – nicht auslesen
    pdf = thumbnails._office_to_pdf(path)
    if pdf:
        return _pdf(str(pdf))
    text = formats.iwork_text(path)
    if not text and (prev := formats.iwork_preview(path)):
        # Neuere iWork-Dateien ohne LibreOffice: gespeicherte Vorschau (erste Seite) lesen – als PDF oder per OCR
        import tempfile

        from . import ocr
        data, ext = prev
        with tempfile.TemporaryDirectory() as tmp:
            f = Path(tmp) / f"preview{ext}"
            f.write_bytes(data)
            if ext == ".pdf":
                return _pdf(str(f))
            if ocr.available():
                try:
                    text = "[Vorschau, per OCR gelesen] " + ocr.ocr_image(str(f))
                except Exception:
                    text = ""
    return text, None


def _is_archive(path: str) -> bool:
    from . import formats
    return formats.is_archive(path)


def extract_text(path: str, ext: str) -> tuple[str, int | None]:
    """-> (text, page_count). Fehler werden nach oben gereicht."""
    if ext == ".pdf":
        t, n = _pdf(path)
    elif ext in (".pptx", ".ppsx"):
        t, n = _pptx(path)
    elif ext == ".docx":
        t, n = _docx(path)
    elif ext in (".xlsx", ".xlsm"):
        t, n = _xlsx(path)
    elif ext == ".ipynb":
        t, n = _ipynb(path)
    elif ext in (".odt", ".ods", ".odp"):
        t, n = _odf(path)
    elif ext in (".eml", ".msg"):
        from . import mail
        m = mail.parse(path)
        t, n = mail.to_text(m), None
    elif ext in VIDEO_EXTS_:
        from . import video
        t, n = video.describe(video.probe(path)), None
    elif ext in IMAGE_LIKE:
        from . import formats
        t = formats.describe_photo(formats.photo_meta(path), ext)
        n = formats.page_count(path) if ext in formats.TIFF_EXTS else None
    elif ext in (".key", ".pages", ".numbers"):
        t, n = _iwork(path)
    elif ext in AUDIO_EXTS_:
        from . import formats
        t, n = formats.describe_audio(formats.audio_meta(path)), None
    elif _is_archive(path):
        from . import formats
        t, n = formats.archive_text(path), None
    elif ext in TEXT_EXTS:
        with open(path, "rb") as fh:
            raw = fh.read(MAX_CHARS * 2)
        for enc in ("utf-8", "cp1252", "latin-1"):
            try:
                t = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        n = None
    else:
        return "", None
    return " ".join(t.split())[:MAX_CHARS] if t else "", n

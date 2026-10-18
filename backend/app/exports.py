"""Exporte: Dateiliste, Vorschlag, Kopierprotokoll, Statistik, Duplikate, ähnliche Dateien, Änderungen, Programme
als Excel (.xlsx) oder CSV – und der Prüfbericht zu einem Kopiervorgang als PDF.

- CSV: UTF-8 mit BOM und Semikolon, öffnet sich in deutschem Excel direkt richtig.
- Excel: eine Tabelle je Blatt mit Kopfzeile, Filter und fixierter erster Zeile; große Mengen im Streaming-Modus.
- Schutz vor Formel-Injektion: Zellen, die mit = + - @ beginnen, werden als Text geschrieben.
- Exporte lesen nur aus der Datenbank. Einzige Schreibstelle ist der Prüfbericht im Zielordner (safety.safe_write).
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from dataclasses import dataclass, field
from datetime import datetime
from typing import Callable, Iterable, Iterator

from . import db
from .config import get_settings

MAX_XLSX_ROWS = 1_048_000  # Excel-Grenze je Blatt (1.048.576)


@dataclass
class Sheet:
    title: str
    headers: list[str]
    rows: Iterable[list]
    widths: list[int] = field(default_factory=list)


# ---------------------------------------------------------------- Hilfen

def _dt(v) -> datetime | None:
    if not v:
        return None
    try:
        d = datetime.fromisoformat(str(v))
        return d.astimezone().replace(tzinfo=None) if d.tzinfo else d
    except ValueError:
        return None


def fmt_bytes(n) -> str:
    v = float(n or 0)
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if v < 1024 or unit == "TB":
            return f"{v:.0f} {unit}" if unit == "B" else f"{v:.1f} {unit}".replace(".", ",")
        v /= 1024
    return ""


def _yn(v) -> str:
    return "ja" if v else "nein"


def _labels(s: str | None) -> str:
    return ", ".join(sorted(set(s.split("|")))) if s else ""


def _iter(sql: str, args: Iterable = ()) -> Iterator[dict]:
    """Zeilen stückweise lesen (auch Millionen Einträge ohne großen Speicher)."""
    conn = db._open(str(get_settings().db_path))  # eigene Verbindung: Streaming-Antworten laufen in wechselnden Threads
    try:
        cur = conn.execute(sql, list(args))
        while chunk := cur.fetchmany(2000):
            for r in chunk:
                yield dict(r)
    finally:
        conn.close()


# ---------------------------------------------------------------- Datensätze

FILE_COLS = """f.id, f.name, c.name computer, v.name volume, f.path, f.file_type, f.ext, f.size, f.mtime, f.ctime,
  f.group_name, f.category, f.importance, f.summary, f.suggested_path, f.final_path, f.review_status, f.dup_count,
  f.is_primary, f.sim_kind, f.sim_best, f.online_only, f.status, f.page_count, f.sha256, f.missing,
  (SELECT group_concat(label, '|') FROM labels l WHERE l.file_id=f.id) labels"""
FILE_FROM = "FROM files f JOIN computers c ON c.id=f.computer_id JOIN volumes v ON v.id=f.volume_id"
SIM_ROLE = {1: "empfohlen", 2: "gleicher Inhalt, anderes Format", 0: "ähnlich/älter"}


def files_sheet(where: str = "f.missing=0", args: Iterable = (), order: str = "f.pinned DESC, f.importance DESC, f.id") -> Sheet:
    headers = ["ID", "Name", "Computer", "Festplatte", "Pfad", "Dateityp", "Endung", "Größe (Bytes)", "Größe", "Geändert",
               "Erstellt", "Bereich", "Kategorie", "Wichtigkeit", "Labels", "Zusammenfassung", "Seiten", "Neuer Pfad (Vorschlag)",
               "Neuer Pfad (bestätigt)", "Prüfstatus", "Duplikate", "Original", "Ähnlich", "Rolle", "Nur online", "Status", "SHA-256"]

    def rows():
        for r in _iter(f"SELECT {FILE_COLS} {FILE_FROM} WHERE {where} ORDER BY {order}", args):
            yield [r["id"], r["name"], r["computer"], r["volume"], r["path"], r["file_type"], r["ext"], r["size"],
                   fmt_bytes(r["size"]), _dt(r["mtime"]), _dt(r["ctime"]), r["group_name"], r["category"], r["importance"],
                   _labels(r["labels"]), r["summary"], r["page_count"], r["suggested_path"], r["final_path"], r["review_status"],
                   r["dup_count"] if (r["dup_count"] or 1) > 1 else None, _yn(r["is_primary"]) if (r["dup_count"] or 1) > 1 else None,
                   r["sim_kind"], SIM_ROLE.get(r["sim_best"]) if r["sim_kind"] else None, _yn(r["online_only"]),
                   "nicht mehr gefunden" if r["missing"] else r["status"], r["sha256"]]
    return Sheet("Dateien", headers, rows(), [8, 36, 16, 18, 60, 14, 8, 14, 10, 17, 17, 11, 26, 11, 30, 50, 7, 50, 40, 12, 10, 9, 11, 16, 10, 12, 20])


def proposal_sheets(include_duplicates: bool = False) -> list[Sheet]:
    p = "COALESCE(f.final_path, f.suggested_path)"
    where = "f.missing=0 AND f.suggested_path IS NOT NULL" + ("" if include_duplicates else " AND f.is_primary=1")
    headers = ["Prüfstatus", "Bereich", "Kategorie", "Neuer Ordner", "Neuer Name", "Vom Nutzer geändert", "Wichtigkeit",
               "Quelle: Computer", "Quelle: Festplatte", "Quelle: Pfad", "Größe (Bytes)", "Größe", "Geändert", "Duplikate",
               "Ähnlich", "Rolle", "Nur online", "Labels", "ID"]

    def rows():
        for r in _iter(f"SELECT {FILE_COLS}, {p} target {FILE_FROM} WHERE {where} ORDER BY {p}", ()):
            folder, _, name = (r["target"] or "").rpartition("/")
            yield [r["review_status"], r["group_name"], r["category"], folder, name, _yn(r["final_path"]), r["importance"],
                   r["computer"], r["volume"], r["path"], r["size"], fmt_bytes(r["size"]), _dt(r["mtime"]),
                   r["dup_count"] if (r["dup_count"] or 1) > 1 else None, r["sim_kind"],
                   SIM_ROLE.get(r["sim_best"]) if r["sim_kind"] else None, _yn(r["online_only"]), _labels(r["labels"]), r["id"]]

    folders = db.query(f"""SELECT rtrim(rtrim({p}, replace({p}, '/', '')), '/') folder, COUNT(*) n, COALESCE(SUM(size),0) bytes,
                           SUM(review_status='freigegeben') ok, SUM(review_status='abgelehnt') no, SUM(review_status='offen') open
                           FROM files f WHERE {where} GROUP BY 1 ORDER BY 1""")
    summary = db.query(f"""SELECT group_name, COUNT(*) n, COALESCE(SUM(size),0) bytes, SUM(review_status='freigegeben') ok,
                           SUM(review_status='abgelehnt') no, SUM(review_status='offen') open FROM files f WHERE {where}
                           GROUP BY 1 ORDER BY 1""")
    return [
        Sheet("Vorschlag", headers, rows(), [12, 11, 26, 50, 36, 10, 11, 16, 18, 60, 14, 10, 17, 10, 11, 18, 10, 30, 8]),
        Sheet("Ordnerstruktur", ["Neuer Ordner", "Dateien", "Größe (Bytes)", "Größe", "freigegeben", "abgelehnt", "offen"],
              [[f["folder"], f["n"], f["bytes"], fmt_bytes(f["bytes"]), f["ok"], f["no"], f["open"]] for f in folders],
              [60, 10, 14, 10, 12, 10, 8]),
        Sheet("Je Bereich", ["Bereich", "Dateien", "Größe (Bytes)", "Größe", "freigegeben", "abgelehnt", "offen"],
              [[s["group_name"], s["n"], s["bytes"], fmt_bytes(s["bytes"]), s["ok"], s["no"], s["open"]] for s in summary],
              [16, 10, 14, 10, 12, 10, 8]),
    ]


def copy_job(job_id: int | None) -> dict | None:
    if job_id:
        return db.one("SELECT * FROM jobs WHERE id=? AND kind='copy'", (job_id,))
    return db.one("SELECT * FROM jobs WHERE kind='copy' ORDER BY id DESC LIMIT 1")


def copy_ops_sheet(job_id: int) -> Sheet:
    headers = ["Status", "Datei", "Bereich", "Kategorie", "Quelle (Original, unverändert)", "Ziel (Kopie)", "Größe (Bytes)",
               "Größe", "SHA-256 geprüft", "Computer", "Meldung", "Erledigt", "SHA-256"]

    def rows():
        for r in _iter("""SELECT o.*, f.name, f.group_name, f.category, f.sha256 FROM copy_ops o JOIN files f ON f.id=o.file_id
                          WHERE o.job_id=? ORDER BY o.id""", (job_id,)):
            yield [r["status"], r["name"], r["group_name"], r["category"], r["src"], r["dst"], r["bytes"], fmt_bytes(r["bytes"]),
                   _yn(r["verified"]), r["computer_id"], r["message"], _dt(r["done_at"] or r["created_at"]), r["sha256"]]
    return Sheet("Kopierprotokoll", headers, rows(), [12, 36, 11, 26, 60, 60, 14, 10, 10, 14, 40, 17, 20])


def stats_sheets() -> list[Sheet]:
    from .api.pipeline import _stats

    s = _stats()
    t = s["totals"]
    k = s.get("llm_kosten") or {}
    overview = [
        ["Dateien erfasst", t["files"]], ["Gesamtgröße", fmt_bytes(t["bytes"])], ["Analysiert", t["analyzed"]],
        ["Fehler bei der Analyse", t["errors"]], ["OCR durchgeführt", t["ocr"]], ["Nur online (Cloud)", t["online_only"]],
        ["Duplikatgruppen", t["dup_groups"]], ["Überzählige Duplikate", t["dup_files"]], ["Davon belegt", fmt_bytes(t["dup_bytes"])],
        ["Gruppen ähnlicher Dateien", t.get("sim_groups")], ["Ältere Versionen", t.get("older_versions")],
        ["Freigegeben", t["approved"]], ["Abgelehnt", t["rejected"]],
        ["Kopien erstellt", (s["copies"] or {}).get("copied")], ["Kopiert (Größe)", fmt_bytes((s["copies"] or {}).get("bytes"))],
        ["Miniaturen", s["thumbnails"]], ["LLM", s["llm_provider"]], ["LLM-Kosten gesamt (USD)", k.get("usd")],
        ["Zielordner", s["target_root"]], ["Erstellt am", datetime.now().replace(microsecond=0)],
    ]

    def grp(title, rows, key="k", extra=()):
        return Sheet(title, ["Wert", "Dateien", "Größe (Bytes)", "Größe", *[e[0] for e in extra]],
                     [[r[key], r["n"], r["bytes"], fmt_bytes(r["bytes"]), *[r.get(e[1]) for e in extra]] for r in rows],
                     [34, 10, 14, 10, 12])

    return [
        Sheet("Übersicht", ["Kennzahl", "Wert"], overview, [32, 50]),
        grp("Je Bereich", s["by_group"]),
        grp("Je Kategorie", s["by_category"], extra=(("Bereich", "grp"), ("Ø Wichtigkeit", "imp"))),
        grp("Je Dateityp", s["by_type"]),
        grp("Je Jahr", s["by_year"]),
        Sheet("Je Computer", ["Computer", "Betriebssystem", "Anbindung", "Dateien", "Größe (Bytes)", "Größe", "Analysiert", "Zuletzt gesehen"],
              [[c["name"], c["os"], c["mode"], c["n"], c["bytes"], fmt_bytes(c["bytes"]), c["analyzed"], _dt(c["last_seen"])] for c in s["by_computer"]],
              [22, 16, 10, 10, 14, 10, 10, 17]),
        Sheet("Je Festplatte", ["Festplatte", "Computer", "Art", "Pfad", "Dateien", "Größe (Bytes)", "Größe", "Kapazität", "Frei", "Zuletzt gelesen"],
              [[v["name"], v["computer_id"], v["kind"], v["path"], v["n"], v["bytes"], fmt_bytes(v["bytes"]),
                fmt_bytes(v["total_bytes"]) if v["total_bytes"] else None, fmt_bytes(v["free_bytes"]) if v["free_bytes"] else None,
                _dt(v["last_scan"])] for v in s["by_volume"]], [24, 14, 10, 50, 10, 14, 10, 11, 11, 17]),
        Sheet("Wichtigkeit", ["Wichtigkeit ab", "Dateien"], [[r["k"], r["n"]] for r in s["importance"]], [16, 10]),
    ]


def duplicates_sheet() -> Sheet:
    headers = ["Gruppe", "Anzahl", "Original", "Name", "Computer", "Festplatte", "Pfad", "Größe (Bytes)", "Größe", "Geändert", "Prüfsumme"]

    def rows():
        g = 0
        last = None
        for r in _iter(f"""SELECT {FILE_COLS}, f.dup_group {FILE_FROM} WHERE f.missing=0 AND f.dup_count>1
                           ORDER BY f.dup_count DESC, f.dup_group, f.is_primary DESC, f.id""", ()):
            if r["dup_group"] != last:
                g, last = g + 1, r["dup_group"]
            yield [g, r["dup_count"], _yn(r["is_primary"]), r["name"], r["computer"], r["volume"], r["path"], r["size"],
                   fmt_bytes(r["size"]), _dt(r["mtime"]), r["dup_group"]]
    return Sheet("Duplikate", headers, rows(), [8, 8, 9, 36, 16, 18, 60, 14, 10, 17, 20])


def similar_sheet() -> Sheet:
    headers = ["Gruppe", "Art", "Anzahl", "Rolle", "Name", "Computer", "Festplatte", "Pfad", "Größe (Bytes)", "Größe", "Geändert", "Bereich", "Kategorie"]

    def rows():
        for r in _iter(f"""SELECT {FILE_COLS}, f.sim_group, f.sim_count {FILE_FROM} WHERE f.missing=0 AND f.sim_group IS NOT NULL
                           ORDER BY f.sim_kind, f.sim_group, f.sim_best=1 DESC, f.mtime DESC""", ()):
            yield [r["sim_group"], r["sim_kind"], r["sim_count"], SIM_ROLE.get(r["sim_best"]), r["name"], r["computer"], r["volume"],
                   r["path"], r["size"], fmt_bytes(r["size"]), _dt(r["mtime"]), r["group_name"], r["category"]]
    return Sheet("Ähnliche Dateien", headers, rows(), [8, 11, 8, 18, 36, 16, 18, 60, 14, 10, 17, 11, 26])


def changes_sheet() -> Sheet:
    headers = ["Zeitpunkt", "Änderung", "Name", "Computer", "Festplatte", "Pfad", "Details", "Prüfstatus"]
    kinds = {"neu": "neu", "geändert": "geändert", "fehlt": "nicht mehr gefunden", "wieder_da": "wieder da"}

    def rows():
        for r in _iter("""SELECT e.ts, e.kind, e.detail, f.name, f.path, f.review_status, c.name computer, v.name volume
                          FROM file_events e JOIN files f ON f.id=e.file_id JOIN computers c ON c.id=f.computer_id
                          JOIN volumes v ON v.id=f.volume_id ORDER BY e.id DESC""", ()):
            yield [_dt(r["ts"]), kinds.get(r["kind"], r["kind"]), r["name"], r["computer"], r["volume"], r["path"], r["detail"], r["review_status"]]
    return Sheet("Änderungen", headers, rows(), [17, 18, 36, 16, 18, 60, 50, 12])


def programs_sheet() -> Sheet:
    headers = ["Computer", "Name", "Version", "Hersteller", "Art", "Kategorie", "Ort"]
    rows = ([r["computer"], r["name"], r["version"], r["publisher"], r["kind"], r["category"], r["location"]]
            for r in _iter("""SELECT COALESCE(c.name, p.computer_id) computer, p.name, p.version, p.publisher, p.kind, p.category, p.location
                              FROM programs p LEFT JOIN computers c ON c.id=p.computer_id ORDER BY 1, p.kind, p.name COLLATE NOCASE""", ()))
    return Sheet("Programme & LLMs", headers, rows, [20, 40, 14, 24, 12, 18, 50])


# ---------------------------------------------------------------- Ausgabe

def _safe_text(v):
    """CSV: Formel-Injektion verhindern (Dateinamen wie „=HYPERLINK(…)“)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def to_csv(sheet: Sheet) -> Iterator[bytes]:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    yield "﻿".encode("utf-8")
    w.writerow(sheet.headers)
    n = 0
    for row in sheet.rows:
        w.writerow(["" if v is None else v.strftime("%d.%m.%Y %H:%M") if isinstance(v, datetime) else _safe_text(v) for v in row])
        n += 1
        if n % 1000 == 0:
            yield buf.getvalue().encode("utf-8")
            buf.seek(0)
            buf.truncate()
    yield buf.getvalue().encode("utf-8")


def to_xlsx(sheets: list[Sheet], title: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F46E5")
    for sh in sheets:
        ws = wb.create_sheet(sh.title[:31])
        for i, w in enumerate(sh.widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        head = []
        for h in sh.headers:
            c = WriteOnlyCell(ws, value=h)
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center")
            head.append(c)
        ws.append(head)
        n = 0
        for row in sh.rows:
            n += 1
            if n > MAX_XLSX_ROWS:
                ws.append([f"… weitere Zeilen – Excel-Grenze erreicht, bitte CSV-Export nutzen"])
                break
            cells = []
            for v in row:
                c = WriteOnlyCell(ws, value=v)
                if isinstance(v, str) and v.startswith("="):
                    c.data_type = "s"  # als Text, nicht als Formel
                elif isinstance(v, datetime):
                    c.number_format = "DD.MM.YYYY HH:MM"
                cells.append(c)
            ws.append(cells)
        ws.auto_filter.ref = f"A1:{get_column_letter(max(len(sh.headers), 1))}{max(n, 0) + 1}"
    wb.properties.title = title or "Future Dateiverwaltung"
    wb.properties.creator = "Future Dateiverwaltung"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- Prüfbericht (PDF)

REPORT_CSS = """
* { font-family: sans-serif; }
body { font-size: 9pt; color: #0f172a; }
h1 { font-size: 17pt; color: #312e81; margin: 0 0 2pt 0; }
h2 { font-size: 11.5pt; color: #312e81; margin: 12pt 0 4pt 0; border-bottom: 1px solid #c7d2fe; }
p { margin: 2pt 0; }
.muted { color: #64748b; }
table { border-collapse: collapse; width: 100%; }
th { background-color: #eef0ff; text-align: left; padding: 2pt 4pt; font-size: 8pt; }
td { padding: 2pt 4pt; border-bottom: 1px solid #e2e8f0; font-size: 8pt; vertical-align: top; }
.ok { color: #047857; font-weight: bold; }
.bad { color: #b91c1c; font-weight: bold; }
.box { background-color: #ecfdf5; border: 1px solid #a7f3d0; padding: 5pt; }
"""
PDF_MAX_ROWS = 3000


def _h(v) -> str:
    from html import escape
    return escape("" if v is None else str(v))


def _path(v, root: str = "") -> str:
    """Pfad für das PDF: relativ zum Zielordner und mit Umbruchstellen nach / bzw. \\ (sonst läuft die Tabelle über)."""
    v = "" if v is None else str(v)
    if root and v.startswith(root):
        v = v[len(root):].lstrip("/\\")
    return _h(v).replace("/", "/\u200b").replace("\\", "\\\u200b").replace("_", "_\u200b")


def report_data(job: dict) -> dict:
    ops = db.query("""SELECT o.*, f.name, f.group_name, f.category FROM copy_ops o JOIN files f ON f.id=o.file_id
                      WHERE o.job_id=? ORDER BY o.status, o.id""", (job["id"],))
    stats = json.loads(job["stats"]) if job.get("stats") else {}
    digest = hashlib.sha256(json.dumps([[o["id"], o["src"], o["dst"], o["status"], o["bytes"], o["verified"]] for o in ops],
                                       ensure_ascii=False).encode()).hexdigest()
    return {"job": job, "stats": stats, "ops": ops, "digest": digest}


def report_pdf(job: dict) -> bytes:
    import pymupdf

    d = report_data(job)
    st, ops = d["stats"], d["ops"]
    verified = sum(1 for o in ops if o["status"] == "kopiert" and o["verified"])
    copied = sum(1 for o in ops if o["status"] == "kopiert")
    errors = [o for o in ops if o["status"] == "fehler"]
    dry = bool(st.get("trockenlauf"))
    kpis = [("Geplant", st.get("geplant")), ("Kopiert", st.get("kopiert")), ("Übersprungen (bereits identisch)", st.get("übersprungen")),
            ("Fehler", st.get("fehler")), ("An Agenten übergeben", st.get("an_agent_übergeben")),
            ("Nur online, nicht heruntergeladen", st.get("nur_online")), ("Serverseitig in der Cloud kopiert", st.get("cloud_serverseitig")),
            ("Ältere Versionen ausgelassen", st.get("ältere_versionen_ausgelassen")), ("Datenmenge kopiert", fmt_bytes(st.get("bytes")))]
    groups = st.get("je_gruppe") or {}
    gbytes = st.get("je_gruppe_bytes") or {}
    html = [f"<h1>Prüfbericht Kopiervorgang #{job['id']}</h1>",
            f"<p class='muted'>Future Dateiverwaltung · erstellt {datetime.now():%d.%m.%Y %H:%M} · Lauf gestartet "
            f"{_dt(job['started_at']):%d.%m.%Y %H:%M} · Status: {_h(job['status'])}{' · TROCKENLAUF – es wurde nichts kopiert' if dry else ''}</p>",
            f"<p><b>Ziel:</b> {_path(st.get('ziel') or get_settings().target_root)}</p>",
            "<h2>Zusagen</h2><div class='box'>",
            "<p>Die Originale wurden nur gelesen – nicht gelöscht, verschoben, umbenannt oder verändert.</p>",
            "<p>Im Ziel wurde nichts überschrieben; bei Namenskonflikten wurde „Name (2)“ angelegt.</p>",
            "<p>Kopiert wurden ausschließlich Dateien, die zuvor manuell freigegeben wurden.</p>",
            f"<p>Jede Kopie wurde per SHA-256 mit dem Original verglichen: <span class='{'ok' if verified == copied else 'bad'}'>"
            f"{verified} von {copied} Kopien bestätigt</span>.</p></div>",
            "<h2>Statistik</h2><table>"]
    html += [f"<tr><td>{_h(k)}</td><td><b>{_h(v if v is not None else 0)}</b></td></tr>" for k, v in kpis]
    html.append("</table>")
    if groups:
        html.append("<h2>Je Bereich</h2><table><tr><th>Bereich</th><th>Dateien</th><th>Größe</th></tr>")
        html += [f"<tr><td>{_h(g)}</td><td>{n}</td><td>{fmt_bytes(gbytes.get(g))}</td></tr>" for g, n in sorted(groups.items())]
        html.append("</table>")
    if errors:
        html.append(f"<h2>Fehler ({len(errors)})</h2><table><tr><th>Datei</th><th>Quelle</th><th>Meldung</th></tr>")
        html += [f"<tr><td>{_h(o['name'])}</td><td>{_path(o['src'])}</td><td class='bad'>{_h(o['message'])}</td></tr>" for o in errors]
        html.append("</table>")
    shown = [o for o in ops if o["status"] != "fehler"][:PDF_MAX_ROWS]
    html.append(f"<h2>Einzelnachweis ({len(ops)} Einträge)</h2>")
    if len(ops) - len(errors) > len(shown):
        html.append(f"<p class='muted'>Gezeigt werden die ersten {len(shown)} Einträge – vollständig im Excel-Protokoll.</p>")
    root = str(st.get("ziel") or get_settings().target_root)
    html.append("<table><tr><th style='width:11%'>Status</th><th style='width:43%'>Datei / Quelle (Original)</th>"
                "<th style='width:32%'>Ziel (im Zielordner)</th><th style='width:7%'>Größe</th><th style='width:7%'>SHA-256</th></tr>")
    html += [f"<tr><td>{_h(o['status'])}</td><td><b>{_h(o['name'])}</b><br/><span class='muted'>{_path(o['src'])}</span></td>"
             f"<td>{_path(o['dst'], root)}</td><td>{fmt_bytes(o['bytes'])}</td><td>{'geprüft' if o['verified'] else '–'}</td></tr>"
             for o in shown]
    html.append("</table>")
    html.append(f"<h2>Prüfsumme dieses Protokolls</h2><p class='muted'>SHA-256 über alle Einträge (ID, Quelle, Ziel, Status, Größe, "
                f"Prüfung) – zum Nachweis, dass das Protokoll nicht verändert wurde:</p><p><b>{d['digest']}</b></p>")

    buf = io.BytesIO()
    story = pymupdf.Story(html="".join(html), user_css=REPORT_CSS)
    writer = pymupdf.DocumentWriter(buf)
    page = pymupdf.paper_rect("a4")
    where = page + (42, 48, -42, -52)
    more = True
    while more:
        dev = writer.begin_page(page)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    doc = pymupdf.open("pdf", buf.getvalue())
    for i, pg in enumerate(doc, 1):  # Fußzeile mit Seitenzahl
        pg.insert_text((42, page.height - 28), f"Future Dateiverwaltung · Prüfbericht Kopiervorgang #{job['id']} · Seite {i} von {doc.page_count}",
                       fontsize=7.5, color=(0.45, 0.5, 0.6))
    doc.set_metadata({"title": f"Prüfbericht Kopiervorgang #{job['id']}", "author": "Future Dateiverwaltung"})
    return doc.tobytes(garbage=3, deflate=True)


def save_report_to_target(job: dict) -> list[str]:
    """Prüfbericht (PDF) und Protokoll (Excel) im Zielordner unter „_Prüfberichte“ ablegen – nie überschreiben."""
    from .safety import safe_write

    stamp = (_dt(job.get("started_at")) or datetime.now()).strftime("%Y-%m-%d_%H%M")
    base = f"_Prüfberichte/Kopie_Lauf_{job['id']}_{stamp}"
    pdf = safe_write(base + ".pdf", report_pdf(job))
    xlsx = safe_write(base + ".xlsx", to_xlsx([copy_ops_sheet(job["id"])], f"Kopierprotokoll #{job['id']}"))
    return [str(pdf), str(xlsx)]


# ---------------------------------------------------------------- Register

EXPORTS: dict[str, tuple[str, Callable[..., list[Sheet]]]] = {
    "statistik": ("Statistik", lambda **k: stats_sheets()),
    "vorschlag": ("Vorschlag", lambda include_duplicates=False, **k: proposal_sheets(include_duplicates)),
    "duplikate": ("Duplikate", lambda **k: [duplicates_sheet()]),
    "aehnliche": ("Ähnliche Dateien", lambda **k: [similar_sheet()]),
    "aenderungen": ("Änderungen", lambda **k: [changes_sheet()]),
    "programme": ("Programme & LLMs", lambda **k: [programs_sheet()]),
}

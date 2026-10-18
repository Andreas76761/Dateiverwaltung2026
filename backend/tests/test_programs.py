"""Programme & LLMs organisieren: Normalisierung, Inventar-Erfassung, Vergleich, Speicher, Empfehlungen."""
import io
import json
import sys
from pathlib import Path

import pytest
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.skills import software  # noqa: E402

GB = 1024 ** 3


def test_model_and_program_names():
    same = [("llama3.1:8b", "Ollama"), ("Meta-Llama-3.1-8B-Instruct-Q4_K_M.gguf", None), ("meta-llama/Llama-3.1-8B-Instruct", "Hugging Face")]
    assert len({software.parse_model(n, t)["key"] for n, t in same}) == 1
    m = software.parse_model("Meta-Llama-3.1-8B-Instruct-Q8_0.gguf")
    assert m["quant"] == "Q8_0" and m["format"] == "GGUF" and m["label"] == "llama 3.1 · 8B · instruct"
    assert software.parse_model("qwen2.5-coder:7b", "Ollama")["key"] == software.parse_model("Qwen/Qwen2.5-Coder-7B-Instruct-GGUF")["key"]
    assert software.parse_model("llama3.1:70b", "Ollama")["key"] != software.parse_model("llama3.1:8b", "Ollama")["key"]
    assert software.parse_model("mixtral:8x7b", "Ollama")["params"] == "8X7B"
    assert software.program_key("Python 3.12.1 (64-bit)")[0] == software.program_key("python3")[0] == "python"
    assert software.program_key("Microsoft Visual Studio Code (User)")[1] == "VS Code"
    assert software.program_key("7-Zip 23.01 (x64)")[1] == "7-Zip"
    assert software.version_tuple("0.5.7") > software.version_tuple("0.3.12")
    assert "Modell-Runner" in software.roles("ollama") and "Coding-Assistent" in software.roles("claude code")


def test_ollama_manifests(tmp_path, monkeypatch):
    from app.agents import inventory_agent
    man = tmp_path / "manifests" / "registry.ollama.ai" / "library" / "llama3.1" / "8b"
    man.parent.mkdir(parents=True)
    man.write_text(json.dumps({"layers": [
        {"mediaType": "application/vnd.ollama.image.model", "digest": "sha256:abc", "size": 4_900_000_000},
        {"mediaType": "application/vnd.ollama.image.template", "digest": "sha256:t", "size": 1_000}]}), encoding="utf-8")
    monkeypatch.setenv("OLLAMA_MODELS", str(tmp_path))
    ms = inventory_agent.ollama_models()
    assert ms == [{"name": "llama3.1:8b", "kind": "llm-modell", "category": "Ollama", "location": str(man),
                   "size": 4_900_001_000, "source": "Ollama-Manifest", "fingerprint": "sha256:abc"}]


@pytest.fixture()
def client(tmp_path, monkeypatch):
    cfg = {"server": {"data_dir": str(tmp_path / "data")}, "target": {"root": str(tmp_path / "ziel")},
           "computers": [{"id": "ws", "name": "Workstation", "volumes": []},
                         {"id": "lap", "name": "Laptop", "mode": "agent", "volumes": []}]}
    (tmp_path / "fdv.yaml").write_text(yaml.safe_dump(cfg), encoding="utf-8")
    monkeypatch.setenv("FDV_CONFIG", str(tmp_path / "fdv.yaml"))
    from app.config import get_settings
    get_settings.cache_clear()
    from fastapi.testclient import TestClient

    from app.main import app
    with TestClient(app) as c:
        yield c


def model(name, size, tool, fp, loc):
    return {"name": name, "kind": "llm-modell", "category": tool, "size": size, "fingerprint": fp, "location": loc}


def post_inventory(c, cid, items):
    assert c.post("/api/agents/inventory", json={"computer_id": cid, "items": items}).status_code == 200


def test_compare_storage_recommendations(client):
    c = client
    post_inventory(c, "ws", [
        {"name": "Ollama", "version": "0.5.7", "kind": "llm-tool", "location": "/usr/bin/ollama"},
        {"name": "LM Studio", "version": "0.3.5", "kind": "llm-tool"},
        {"name": "GPT4All", "version": "3.4", "kind": "llm-tool"},
        {"name": "Python 3.12.7 (64-bit)", "version": "3.12.7"},
        model("llama3.1:8b", 5 * GB, "Ollama", "sha256:l8", "~/.ollama/l8"),
        model("Meta-Llama-3.1-8B-Instruct-Q4_K_M.gguf", 5 * GB, "LM Studio", "q:q4", "~/.lmstudio/q4.gguf"),
        model("Meta-Llama-3.1-8B-Instruct-Q4_K_M.gguf", 5 * GB, "GPT4All", "q:q4", "~/gpt4all/q4.gguf"),
        model("Meta-Llama-3.1-8B-Instruct-Q8_0.gguf", 8 * GB, "LM Studio", "q:q8", "~/.lmstudio/q8.gguf"),
        model("llama2:7b", 4 * GB, "Ollama", "sha256:l2", "~/.ollama/l2"),
        {"name": "GPU: RTX 4070", "kind": "hardware", "category": "GPU", "size": 12 * GB},
    ])
    post_inventory(c, "lap", [
        {"name": "Ollama", "version": "0.3.12", "kind": "llm-tool"},
        {"name": "python3", "version": "3.11.4", "kind": "dev-tool"},
        model("llama3.1:8b", 5 * GB, "Ollama", "sha256:l8", "~/.ollama/l8"),
        model("llama3.1:70b", 40 * GB, "Ollama", "sha256:l70", "~/.ollama/l70"),
        {"name": "Arbeitsspeicher (RAM)", "kind": "hardware", "category": "RAM", "size": 16 * GB},
    ])

    cmp = c.get("/api/programs/compare").json()
    rows = {r["display"]: r for r in cmp["rows"]}
    assert rows["Ollama"]["computers"] == 2 and rows["Ollama"]["drift"] and rows["Ollama"]["cells"]["lap"]["outdated"]
    assert rows["Python"]["computers"] == 2  # „Python 3.12.7 (64-bit)“ und „python3“ sind dasselbe
    l8 = rows["llama 3.1 · 8B · instruct"]
    assert l8["computers"] == 2 and l8["cells"]["ws"]["count"] == 4

    st = c.get("/api/programs/storage").json()
    ws = next(x for x in st["computers"] if x["id"] == "ws")
    assert ws["models"] == 27 * GB and ws["vram"] == 12 * GB and st["largest"][0]["size"] == 40 * GB

    recs = c.get("/api/programs/recommendations").json()
    by_type = {}
    for r in recs["items"]:
        by_type.setdefault(r["type"], []).append(r)
    assert by_type["doppelt"][0]["saving"] == 5 * GB                      # identische GGUF in LM Studio + GPT4All
    assert by_type["mehrfach"][0]["saving"] == 5 * GB + 5 * GB            # Ollama + Q4 + Q8 → größte (8 GB) behalten
    assert any(r["type"] == "ueberschneidung" and "Modell-Runner" in r["title"] for r in recs["items"])
    assert {r["title"].split(":")[0] for r in by_type["version"]} >= {"Ollama", "Python"}
    assert by_type["hardware"][0]["items"][0]["name"] == "llama3.1:70b"  # 40 GB bei 16 GB RAM ohne GPU
    assert "Workstation" not in by_type["hardware"][0]["detail"]          # 12 GB GPU reicht auch nicht
    assert by_type["veraltet"][0]["items"][0]["name"] == "llama2:7b"
    assert by_type["zentral"][0]["computers"] == ["Laptop", "Workstation"]
    assert recs["einsparung"] > 0

    key = by_type["doppelt"][0]["key"]
    assert c.post(f"/api/programs/recommendations/{key}", json={"status": "erledigt"}).status_code == 200
    after = c.get("/api/programs/recommendations").json()
    assert next(r for r in after["items"] if r["key"] == key)["status"] == "erledigt"
    assert after["offen"] == recs["offen"] - 1
    assert c.post(f"/api/programs/recommendations/{key}", json={"status": "weg"}).status_code == 400

    # Neues Inventar ohne Llama 2: alter Eintrag wird inaktiv (nicht gelöscht) und fällt aus dem Vergleich
    post_inventory(c, "ws", [{"name": "Ollama", "version": "0.5.7", "kind": "llm-tool", "location": "/usr/bin/ollama"}])
    assert "llama 2 · 7B · instruct" not in {r["display"] for r in c.get("/api/programs/compare").json()["rows"]}
    assert any(p["name"] == "llama2:7b" for p in c.get("/api/programs", params={"include_inactive": True}).json()["items"])

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(c.get("/api/export/programme.xlsx").content))
    assert wb.sheetnames == ["Empfehlungen", "Vergleich", "Speicher", "Inventar"]
    assert c.get("/api/export/programme.csv").content.decode("utf-8-sig").startswith("Computer;Name;Vergleichsname")

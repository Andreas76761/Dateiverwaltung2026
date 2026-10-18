"""Exporte (Excel/CSV) und Prüfbericht (PDF) – inkl. Schutz vor Formel-Injektion und „nie überschreiben“."""
import io
import sys
import time
from pathlib import Path

import pytest
import yaml

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


@pytest.fixture()
def env(tmp_path, monkeypatch):
    src = tmp_path / "pc" / "Dokumente"
    src.mkdir(parents=True)
    (src / "Steuerbescheid_2023.txt").write_text("Finanzamt Steuerbescheid Einkommensteuer 2023 Betrag 1.200,00 EUR", encoding="utf-8")
    (src / "Kopie").mkdir()
    (src / "Kopie" / "Steuerbescheid_2023.txt").write_text("Finanzamt Steuerbescheid Einkommensteuer 2023 Betrag 1.200,00 EUR", encoding="utf-8")
    (src / "=HYPERLINK(1).txt").write_text("Notiz mit gefährlichem Dateinamen", encoding="utf-8")
    (src / "Über_Ärger.md").write_text("System prompt: Du bist ein KI-Agent mit RAG", encoding="utf-8")
    target = tmp_path / "cloud" / "Ziel"
    cfg = {"server": {"data_dir": str(tmp_path / "data")}, "target": {"root": str(target)},
           "analysis": {"ocr": False, "thumbnails": False},
           "computers": [{"id": "pc", "name": "PC", "volumes": [{"id": "pc-c", "name": "C", "path": str(tmp_path / "pc")}]}]}
    (tmp_path / "fdv.yaml").write_text(yaml.safe_dump(cfg), encoding="utf-8")
    monkeypatch.setenv("FDV_CONFIG", str(tmp_path / "fdv.yaml"))
    from app.config import get_settings
    get_settings.cache_clear()
    from fastapi.testclient import TestClient

    from app.main import app
    with TestClient(app) as client:
        for kind in ("scan", "analyze"):
            wait(client, client.post(f"/api/pipeline/{'scan' if kind == 'scan' else 'analyze'}", json={}).json()["job_id"])
        yield client, src, target


def wait(client, job_id):
    for _ in range(300):
        j = client.get(f"/api/jobs/{job_id}").json()
        if j["status"] != "läuft":
            return j
        time.sleep(0.1)
    raise TimeoutError


def xlsx(resp):
    from openpyxl import load_workbook
    assert resp.status_code == 200, resp.text[:300]
    assert "attachment" in resp.headers["content-disposition"]
    return load_workbook(io.BytesIO(resp.content))


def test_file_list_csv_and_xlsx(env):
    client, src, _ = env
    r = client.get("/api/export/dateien.csv", params={"group": "Privat"})
    assert r.status_code == 200 and r.content.startswith("﻿".encode())
    text = r.content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0].startswith("ID;Name;Computer")
    assert all("Privat" in l for l in lines[1:]) and len(lines) == 3  # Steuerbescheid + Duplikat

    wb = xlsx(client.get("/api/export/dateien.xlsx"))
    ws = wb["Dateien"]
    rows = list(ws.values)
    names = [r[1] for r in rows[1:]]
    assert {"Über_Ärger.md", "=HYPERLINK(1).txt"} <= set(names)
    cell = next(c for c in ws["B"] if c.value == "=HYPERLINK(1).txt")
    assert cell.data_type == "s"  # als Text, nicht als Formel
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    # CSV: gefährliche Zelle entschärft
    assert "'=HYPERLINK(1).txt" in client.get("/api/export/dateien.csv").content.decode("utf-8-sig")


def test_other_exports(env):
    client, *_ = env
    wb = xlsx(client.get("/api/export/statistik.xlsx"))
    assert {"Übersicht", "Je Bereich", "Je Kategorie", "Je Computer", "Je Festplatte"} <= set(wb.sheetnames)
    over = dict(list(wb["Übersicht"].values)[1:])
    assert over["Dateien erfasst"] == 4
    wb = xlsx(client.get("/api/export/vorschlag.xlsx"))
    assert wb.sheetnames == ["Vorschlag", "Ordnerstruktur", "Je Bereich"]
    assert len(list(wb["Vorschlag"].values)) == 4  # Kopfzeile + 3 Originale (ohne Duplikat)
    assert len(list(xlsx(client.get("/api/export/vorschlag.xlsx", params={"include_duplicates": True}))["Vorschlag"].values)) == 5
    dup = list(xlsx(client.get("/api/export/duplikate.xlsx"))["Duplikate"].values)
    assert len(dup) == 3 and {r[2] for r in dup[1:]} == {"ja", "nein"}
    for kind in ("aehnliche", "aenderungen", "programme"):
        assert client.get(f"/api/export/{kind}.csv").status_code == 200
    assert client.get("/api/export/unbekannt.xlsx").status_code == 404
    assert client.get("/api/export/statistik.docx").status_code == 400


def test_copy_report(env):
    import pymupdf

    client, src, target = env
    client.post("/api/review", json={"action": "freigeben", "folder": ""})
    j = wait(client, client.post("/api/pipeline/copy", json={}).json()["job_id"])
    assert j["stats"]["kopiert"] == 3
    reports = sorted(p.name for p in (target / "_Prüfberichte").iterdir())
    assert len(reports) == 2 and reports[0].endswith(".pdf") and reports[1].endswith(".xlsx")  # automatisch im Zielordner

    r = client.get(f"/api/export/kopie/{j['id']}/pruefbericht.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    doc = pymupdf.open("pdf", r.content)
    text = "".join(p.get_text() for p in doc)
    assert "Prüfbericht Kopiervorgang" in text and "3 von 3 Kopien bestätigt" in text
    assert "Steuerbescheid_2023.txt" in text and "Seite 1 von" in text

    ops = list(xlsx(client.get(f"/api/export/kopie/{j['id']}.xlsx"))["Kopierprotokoll"].values)
    assert len(ops) == 4 and all(r[0] == "kopiert" and r[8] == "ja" for r in ops[1:])

    # Erneut ablegen überschreibt nie – es entsteht „(2)“
    saved = client.post(f"/api/export/kopie/{j['id']}/pruefbericht-ablegen").json()["dateien"]
    assert all(" (2)" in Path(p).name for p in saved)
    assert len(list((target / "_Prüfberichte").iterdir())) == 4
    # Trockenlauf legt keinen Bericht ab
    wait(client, client.post("/api/pipeline/copy", json={"dry_run": True}).json()["job_id"])
    assert len(list((target / "_Prüfberichte").iterdir())) == 4
    assert client.get("/api/export/kopie/999/pruefbericht.pdf").status_code == 404

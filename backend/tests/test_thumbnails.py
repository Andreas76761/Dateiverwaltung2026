"""Miniaturen aller Seiten für PDF, PPTX, DOCX, XLSX, PNG, JPEG."""
import shutil
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

needs_office = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice fehlt")


@pytest.fixture()
def thumbs(tmp_path, monkeypatch):
    monkeypatch.setenv("FDV_CONFIG", str(tmp_path / "keine.yaml"))
    monkeypatch.setenv("FDV_DATA_DIR", str(tmp_path / "data"))
    from app.config import get_settings
    get_settings.cache_clear()
    from app.skills import thumbnails
    yield thumbnails
    get_settings.cache_clear()


def test_pdf_and_images(thumbs, tmp_path):
    import pymupdf
    from PIL import Image
    doc = pymupdf.open()
    for i in range(3):
        doc.new_page().insert_text((50, 50), f"Seite {i + 1}")
    doc.save(tmp_path / "a.pdf")
    assert len(thumbs.generate(1, str(tmp_path / "a.pdf"), ".pdf")) == 3
    Image.new("RGBA", (300, 200), (0, 0, 255, 128)).save(tmp_path / "b.png")
    Image.new("RGB", (300, 200), "red").save(tmp_path / "c.jpg")
    assert len(thumbs.generate(2, str(tmp_path / "b.png"), ".png")) == 1
    assert len(thumbs.generate(3, str(tmp_path / "c.jpg"), ".jpg")) == 1


@needs_office
def test_docx_all_pages(thumbs, tmp_path):
    import docx
    d = docx.Document()
    for i in range(3):
        if i:
            d.add_page_break()
        d.add_heading(f"Kapitel {i + 1}", 1)
    d.save(tmp_path / "d.docx")
    assert len(thumbs.generate(4, str(tmp_path / "d.docx"), ".docx")) == 3


@needs_office
def test_xlsx_one_thumbnail_per_sheet(thumbs, tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append([f"Spalte {c}" for c in range(30)])  # breit: darf nicht auf mehrere Seiten zerfallen
    long = wb.create_sheet("Lang")
    for i in range(500):
        long.append([f"Artikel {i}", i])
    wb.create_sheet("Leer").append(["x"])
    wb.save(tmp_path / "e.xlsx")
    res = thumbs.generate(5, str(tmp_path / "e.xlsx"), ".xlsx")
    assert len(res) == 3
    assert all(r["width"] == thumbs.THUMB_WIDTH for r in res)
    assert max(r["height"] for r in res) <= thumbs.THUMB_WIDTH * 1.42  # lange Blätter auf A4 begrenzt
